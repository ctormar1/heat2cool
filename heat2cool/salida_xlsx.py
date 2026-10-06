"""Excel de resultados de un modo, con gráficos nativos de Excel y formato condicional."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import parametros as P
from .motor import CALIDAD, INF

CAB = PatternFill("solid", fgColor="1F4E78")
BLANCO = Font(color="FFFFFF", bold=True)
TIT = Font(size=16, bold=True, color="1F4E78")
SUB = Font(size=11, bold=True, color="1F4E78")
COLOR_CAL = {0: "BDD7EE", 1: "C6EFCE", 2: "FFEB9C", 3: "FFC7CE"}
EMOJI_CAL = {0: "◆ genérico", 1: "● dentro de rango", 2: "▲ extrapolado", 3: "✖ fuera de rango"}
VERDE_ROJO = dict(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84",
                  end_type="max", end_color="63BE7B")
ROJO_VERDE = dict(start_type="min", start_color="63BE7B", mid_type="percentile", mid_value=50, mid_color="FFEB84",
                  end_type="max", end_color="F8696B")

NOMBRES = {  # columna -> (cabecera, formato)
    "ranking": ("Ranking", "0"), "cop_bdc1": ("COP BdC 1", "0.00"), "cop_abs": ("COP absorción", "0.000"),
    "cop_r2": ("COP BdC R2", "0.00"), "eer_enf": ("EER enfriadora", "0.00"), "cal_bdc1": ("Calidad BdC 1", "0"),
    "cal_abs": ("Calidad absorción", "0"), "cal_r2": ("Calidad R2", "0"), "n_bdc1": ("Nº BdC 1", "0"),
    "n_abs": ("Nº absorción", "0"), "n_r2": ("Nº BdC R2", "0"), "limitante": ("Qué limita a plena carga", "@"),
    "frio_dis_kw": ("Frío absorción a plena carga (kW)", "0.0"), "cobertura_frio": ("Cobertura de la demanda", "0%"),
    "frio_mwh": ("Frío absorción (MWh/a)", "0.0"), "calor_rec_mwh": ("Calor recuperado (MWh/a)", "0.0"),
    "elec_sis_mwh": ("Electricidad sistema nuevo (MWh/a)", "0.0"),
    "elec_enf_mwh": ("Electricidad enfriadora residual (MWh/a)", "0.0"),
    "elec_base_mwh": ("Electricidad línea base (MWh/a)", "0.0"),
    "ahorro_elec_mwh": ("Ahorro de electricidad (MWh/a)", "0.0"),
    "agua_m3": ("Agua (m3/a)", "#,##0"), "agua_base_m3": ("Agua línea base (m3/a)", "#,##0"),
    "eer_sis": ("EER sistema nuevo", "0.00"), "capex_total": ("CAPEX total (€)", "#,##0"),
    "c_elec": ("Coste electricidad (€/a)", "#,##0"), "c_agua": ("Coste agua (€/a)", "#,##0"),
    "c_mant": ("Mantenimiento (€/a)", "#,##0"), "capex_anual": ("CAPEX anualizado (€/a)", "#,##0"),
    "credito_calor": ("Crédito por calor (€/a)", "#,##0"), "J": ("Coste anual J (€/a)", "#,##0"),
    "J_base": ("Coste anual línea base (€/a)", "#,##0"), "ahorro_neto": ("AHORRO NETO (€/a)", "#,##0"),
    "retorno_anos": ("Retorno simple (años)", "0.0"), "eer_equilibrio": ("EER enfriadora de equilibrio", "0.00"),
    "elec_equilibrio": ("Precio elec. de equilibrio (€/kWh)", "0.000"),
}


def _cab(ws, fila, cols, col0=1):
    for i, c in enumerate(cols):
        cell = ws.cell(fila, col0 + i, c)
        cell.fill, cell.font = CAB, BLANCO
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _anchos(ws, anchos):
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _v(x):
    if x is None:
        return "—"
    if isinstance(x, float) and x == INF:
        return "∞"
    return x


def _tabla(ws, fila, filas_kv):
    for k, v, fmt in filas_kv:
        ws.cell(fila, 1, k)
        c = ws.cell(fila, 2, _v(v))
        if fmt and not isinstance(c.value, str):
            c.number_format = fmt
        fila += 1
    return fila


def _poner(ws, ch, ancla):
    """openpyxl ≥3.1 marca los ejes como borrados y las series invierten color si son negativas: se corrige aquí."""
    for eje in (ch.x_axis, ch.y_axis):
        eje.delete = False
    ch.x_axis.tickLblPos = "low"  # etiquetas de categoría fuera del área aunque haya valores negativos
    for s in ch.series:
        s.invertIfNegative = False
    ws.add_chart(ch, ancla)


def _calidad(cell, nivel):
    cell.value = EMOJI_CAL[nivel]
    cell.fill = PatternFill("solid", fgColor=COLOR_CAL[nivel])


def _hoja_resumen(ws, modo, caso, res, equil):
    rec = modo == "con_recuperacion"
    _anchos(ws, [50, 18, 4, 26, 14, 14])
    ws["A1"] = f"heat2cool · {P.MODOS[modo]}"
    ws["A1"].font = TIT
    ws["A2"] = "Caso de la hoja Entradas. Los importes dependen de los parámetros [SUPUESTO]."
    ws["A2"].font = Font(italic=True, color="808080")
    ws["A4"] = "Indicadores clave"
    ws["A4"].font = SUB
    f = _tabla(ws, 5, [
        ("AHORRO NETO anual frente a línea base (€/a)", res["ahorro_neto"], "#,##0"),
        ("CAPEX total (€)", res["capex_total"], "#,##0"),
        ("Retorno simple (años)", res["retorno_anos"], "0.0"),
        ("Coste anual J sistema (€/a)", res["J"], "#,##0"),
        ("Coste anual línea base (€/a)", res["J_base"], "#,##0"),
        ("Ahorro de electricidad (MWh/a)", res["ahorro_elec_mwh"], "0.0"),
        ("Frío producido por absorción (MWh/a)", res["frio_mwh"], "0.0"),
        ("Calor recuperado útil (MWh/a)", res["calor_rec_mwh"], "0.0"),
        ("EER del sistema nuevo (kWh frío / kWhe)", res["eer_sis"], "0.00"),
        ("Cobertura de la demanda de frío a plena carga", res["cobertura_frio"], "0%"),
        ("Qué limita a plena carga", res["limitante"], None),
        ("EER de la enfriadora existente con el que se empata", equil["eer_equilibrio"], "0.00"),
        ("Precio de electricidad con el que se empata (€/kWh)", equil["elec_equilibrio"], "0.000"),
    ])
    ws["B5"].font = Font(bold=True, size=13, color="00B050" if res["ahorro_neto"] > 0 else "C00000")

    # tabla traspuesta (filas = base / nuevo, columnas = partidas) para que cada partida sea una serie apilada
    ws["D4"] = "Coste anual (€/a)"
    ws["D4"].font = SUB
    partidas = ["Electricidad", "Agua", "Mantenimiento", "CAPEX anualizado", "Crédito por calor (−)"]
    _cab(ws, 5, [""] + partidas, 4)
    for i, (n, vals) in enumerate([("Línea base", [res["c_elec_base"], res["c_agua_base"], 0, 0, 0]),
                                   ("Sistema nuevo", [res["c_elec"], res["c_agua"], res["c_mant"],
                                                      res["capex_anual"], -res["credito_calor"]])]):
        ws.cell(6 + i, 4, n).font = Font(bold=True)
        for j, v in enumerate(vals):
            ws.cell(6 + i, 5 + j, v).number_format = "#,##0"
    for j in range(5, 10):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap = "col", "stacked", 100
    ch.title, ch.y_axis.title = "Coste anual: línea base frente a sistema nuevo", "€/año"
    ch.add_data(Reference(ws, min_col=5, max_col=9, min_row=5, max_row=7), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=4, min_row=6, max_row=7))
    ch.height, ch.width = 9, 15
    _poner(ws, ch, "D10")

    ws.cell(f + 1, 1, "Equipos seleccionados").font = SUB
    _cab(ws, f + 2, ["Equipo", "Modelo"])
    ws.cell(f + 2, 4, "Calidad del dato").font = SUB
    eqs = [("BdC 1", caso["eq.bdc1"], res["cal_bdc1"]), ("Absorción", caso["eq.abs"], res["cal_abs"])]
    if rec:
        eqs.append(("BdC R2", caso["eq.r2"], res["cal_r2"]))
    for i, (n, m, cal) in enumerate(eqs):
        ws.cell(f + 3 + i, 1, n)
        ws.cell(f + 3 + i, 2, m)
        _calidad(ws.cell(f + 3 + i, 4), cal)


def _hoja_balance(ws, caso, d):
    _anchos(ws, [34, 12, 14, 14, 14, 4, 32, 12])
    ws["A1"] = "Balance de circuitos a plena carga (diseño, sin recuperación)"
    ws["A1"].font = TIT
    _cab(ws, 3, ["Circuito", "T ida (°C)", "T retorno (°C)", "Potencia (kW)", "Caudal (m3/h)"])
    circ = [("CPD → evaporador BdC 1", caso["cpd.t_sal"], caso["cpd.t_ret"], d["q_evap"], d["f_cpd"]),
            ("BdC 1 → generador absorción", caso["cal.t_ida"], caso["cal.t_ret"], d["q_gen"], d["f_cal"]),
            ("Agua enfriada (absorción)", caso["frio.t_imp"], caso["frio.t_ret"], d["q_frio"], d["f_frio"]),
            ("Agua de refrigeración (absorción)", caso["rech.t_ent"], caso["rech.t_sal"], d["q_rej"], d["f_rech"])]
    for i, row in enumerate(circ):
        for j, v in enumerate(row):
            c = ws.cell(4 + i, 1 + j, v)
            if j >= 3:
                c.number_format = "0.0"
    _cab(ws, 3, ["Flujo de energía", "kW"], 7)
    flujos = [("Calor del CPD aprovechado", d["q_evap"]), ("Electricidad BdC 1", d["w_bdc1"]),
              ("Calor al generador", d["q_gen"]), ("Frío por absorción", d["q_frio"]),
              ("Frío con enfriadora existente", d["q_enf"]), ("Rechazo absorción", d["q_rej"]),
              ("A torre existente", d["q_torre"]), ("A adiabático nuevo", d["q_adiab"]),
              ("Bombas (kWe)", d["p_bombas"]), ("Ventiladores (kWe)", d["p_vent"])]
    for i, (n, v) in enumerate(flujos):
        ws.cell(4 + i, 7, n)
        ws.cell(4 + i, 8, v).number_format = "0.0"
    ch = BarChart()
    ch.type, ch.title, ch.y_axis.title, ch.legend = "bar", "Flujos de energía a plena carga", "kW", None
    ch.x_axis.scaling.orientation = "maxMin"
    ch.y_axis.crosses = "max"
    ch.add_data(Reference(ws, min_col=8, min_row=4, max_row=3 + len(flujos)), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=7, min_row=4, max_row=3 + len(flujos)))
    ch.height, ch.width = 9, 16
    _poner(ws, ch, "A10")


def _hoja_equipos(ws, modo, caso, res, cat):
    _anchos(ws, [20, 52, 12, 14, 22, 20, 20, 80])
    ws["A1"] = "Equipos: modelo, dimensionado y procedencia del dato"
    ws["A1"].font = TIT
    _cab(ws, 3, ["Equipo", "Modelo", "Nº unidades", "COP en diseño", "Capacidad/ud en diseño (kW)",
                 "Nominal comprada (kW)", "Calidad del dato", "Fuente"])
    filas = [("BdC 1", caso["eq.bdc1"], res["n_bdc1"], res["cop_bdc1"], res["cap_bdc1_kw"], res["nom_bdc1_kw"], res["cal_bdc1"]),
             ("Absorción", caso["eq.abs"], res["n_abs"], res["cop_abs"], res["cap_abs_kw"], res["nom_abs_kw"], res["cal_abs"])]
    if modo == "con_recuperacion":
        filas.append(("BdC R2", caso["eq.r2"], res["n_r2"], res["cop_r2"], res["cap_r2_kw"], res["nom_r2_kw"], res["cal_r2"]))
    for i, (n, m, nu, cop, cap, nom, cal) in enumerate(filas):
        r = 4 + i
        for j, v in enumerate([n, m, nu, cop, "capacidad libre" if cap == INF else cap, nom]):
            c = ws.cell(r, 1 + j, _v(v))
            if j in (3, 4, 5) and not isinstance(c.value, str):
                c.number_format = "0.00" if j == 3 else "0.0"
        _calidad(ws.cell(r, 7), cal)
        ws.cell(r, 8, cat.get(m, {}).get("fuente", ""))
    r = 4 + len(filas)
    ws.cell(r, 1, "Enfriadora existente")
    ws.cell(r, 2, "genérica (EER de la hoja Entradas)")
    ws.cell(r, 4, res["eer_enf"]).number_format = "0.00"
    r += 2
    ws.cell(r, 1, "Leyenda de calidad").font = SUB
    for k, txt in CALIDAD.items():
        r += 1
        _calidad(ws.cell(r, 1), k)
        ws.cell(r, 2, txt)


def _hoja_barrido(wb, rutas, filas):
    ws = wb.create_sheet("Barrido")
    cols = ["ranking"] + rutas + [k for k in NOMBRES if k != "ranking" and k in filas[0]]
    _cab(ws, 1, [NOMBRES[k][0] if k in NOMBRES else f"{P.etiqueta(k)} ({P.unidad(k)})" for k in cols])
    ws.row_dimensions[1].height = 48
    for i, fila in enumerate(filas):
        for j, k in enumerate(cols):
            c = ws.cell(2 + i, 1 + j, _v(fila[k]))
            if k in NOMBRES and not isinstance(c.value, str):
                c.number_format = NOMBRES[k][1]
            if k.startswith("cal_"):
                _calidad(c, fila[k])
    n = len(filas)
    for k in ("ahorro_neto", "eer_sis"):
        letra = get_column_letter(cols.index(k) + 1)
        ws.conditional_formatting.add(f"{letra}2:{letra}{n + 1}", ColorScaleRule(**VERDE_ROJO))
    for j in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
    for k in ("limitante", "cal_bdc1", "cal_abs", "cal_r2"):
        ws.column_dimensions[get_column_letter(cols.index(k) + 1)].width = 26 if k == "limitante" else 17
    ws.freeze_panes = ws.cell(2, 2 + len(rutas))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n + 1}"


def _hoja_mapas(wb, rutas, filas):
    wm = wb.create_sheet("Mapas")
    if len(rutas) == 1:
        x_r = rutas[0]
        orden = sorted(filas, key=lambda f: f[x_r])
        _cab(wm, 1, [P.etiqueta(x_r), "Ahorro neto (€/a)", "EER sistema", "CAPEX (€)"])
        for i, f in enumerate(orden):
            for j, v in enumerate([f[x_r], f["ahorro_neto"], f["eer_sis"], f["capex_total"]]):
                wm.cell(2 + i, 1 + j, v)
        lc = LineChart()
        lc.title = f"Ahorro neto frente a {P.etiqueta(x_r)}"
        lc.add_data(Reference(wm, min_col=2, min_row=1, max_row=1 + len(orden)), titles_from_data=True)
        lc.set_categories(Reference(wm, min_col=1, min_row=2, max_row=1 + len(orden)))
        _poner(wm, lc, "F2")
        return
    # con 3+ variables, el mapa toma el mejor valor sobre el resto de variables
    x_r, y_r = rutas[0], rutas[1]
    xs = sorted({f[x_r] for f in filas})
    ys = sorted({f[y_r] for f in filas})
    nota = "" if len(rutas) == 2 else "  (mejor valor sobre las demás variables)"
    fila0 = 1
    for clave, fmt, mejor_alto in [("ahorro_neto", "#,##0", True), ("eer_sis", "0.00", True),
                                   ("cop_abs", "0.000", True), ("capex_total", "#,##0", False),
                                   ("retorno_anos", "0.0", False)]:
        tab = {}
        for f in filas:
            k = (f[x_r], f[y_r])
            v = f[clave]
            if k not in tab or (v > tab[k] if mejor_alto else v < tab[k]):
                tab[k] = v
        wm.cell(fila0, 1, f"{NOMBRES[clave][0]}  ·  filas: {P.etiqueta(y_r)}  ·  columnas: {P.etiqueta(x_r)}{nota}").font = SUB
        _cab(wm, fila0 + 1, [f"{P.unidad(y_r)} \\ {P.unidad(x_r)}"] + xs)
        for i, yv in enumerate(ys):
            wm.cell(fila0 + 2 + i, 1, yv).font = Font(bold=True)
            for j, xv in enumerate(xs):
                c = wm.cell(fila0 + 2 + i, 2 + j, _v(tab.get((xv, yv))))
                if not isinstance(c.value, str):
                    c.number_format = fmt
        rng = f"B{fila0 + 2}:{get_column_letter(1 + len(xs))}{fila0 + 1 + len(ys)}"
        wm.conditional_formatting.add(rng, ColorScaleRule(**(VERDE_ROJO if mejor_alto else ROJO_VERDE)))
        alto = len(ys) + 4
        if clave == "ahorro_neto":
            lc = LineChart()
            lc.title = f"Ahorro neto frente a {P.etiqueta(x_r)} (una línea por {P.etiqueta(y_r)})"
            lc.y_axis.title, lc.x_axis.title = "€/año", f"{P.etiqueta(x_r)} ({P.unidad(x_r)})"
            lc.add_data(Reference(wm, min_col=1, max_col=1 + len(xs), min_row=fila0 + 2, max_row=fila0 + 1 + len(ys)),
                        from_rows=True, titles_from_data=True)
            lc.set_categories(Reference(wm, min_col=2, max_col=1 + len(xs), min_row=fila0 + 1))
            lc.height, lc.width = 9, 17
            _poner(wm, lc, f"{get_column_letter(len(xs) + 3)}{fila0}")
            alto = max(alto, 20)
        fila0 += alto
    wm.column_dimensions["A"].width = 14


def _hoja_anual(ws, res):
    ws["A1"] = "Funcionamiento por nivel de carga del CPD (potencias medias en kW)"
    ws["A1"].font = TIT
    cl = [("carga", "Carga CPD", "0%"), ("rec", "Recuperación", "@"), ("peso", "Fracción de horas", "0%"),
          ("limitante", "Qué limita", "@"), ("cop_bdc1", "COP BdC 1", "0.00"), ("cop_abs", "COP abs.", "0.000"),
          ("plr_bdc1", "Carga BdC 1", "0%"), ("plr_abs", "Carga abs.", "0%"), ("q_evap", "Calor CPD", "0.0"),
          ("w_bdc1", "Elec. BdC 1", "0.0"), ("q_gen", "Calor generador", "0.0"), ("q_frio", "Frío absorción", "0.0"),
          ("q_enf", "Frío enfriadora", "0.0"), ("p_enf", "Elec. enfriadora", "0.0"), ("q_rej", "Rechazo", "0.0"),
          ("q_r1", "Calor R1", "0.0"), ("q_r2", "Calor R2", "0.0"), ("w_r2", "Elec. R2", "0.0"),
          ("q_torre", "A torre", "0.0"), ("q_adiab", "A adiabático", "0.0"), ("p_bombas", "Bombas", "0.00"),
          ("p_vent", "Ventiladores", "0.00"), ("p_sis", "Elec. sistema neta", "0.0"), ("agua_l_h", "Agua (L/h)", "0")]
    _cab(ws, 3, [x[1] for x in cl])
    ws.row_dimensions[3].height = 32
    cargas = res["cargas"]
    for i, p in enumerate(cargas):
        for j, (k, _, fmt) in enumerate(cl):
            v = ("sí" if p[k] else "no") if k == "rec" else p[k]
            ws.cell(4 + i, 1 + j, v).number_format = fmt
    for j in range(1, len(cl) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 12
    ws.column_dimensions["D"].width = 26
    nc = len(cargas)
    for titulo, columnas, ancla in [("Frío por nivel de carga: absorción + enfriadora", (12, 13), f"A{nc + 6}"),
                                    ("Electricidad por nivel de carga", (10, 14, 18, 21, 22), f"J{nc + 6}")]:
        ch = BarChart()
        ch.type, ch.grouping, ch.overlap = "col", "stacked", 100
        ch.title, ch.y_axis.title = titulo, "kW"
        for col in columnas:
            ch.add_data(Reference(ws, min_col=col, min_row=3, max_row=3 + nc), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, max_col=2, min_row=4, max_row=3 + nc))
        ch.height, ch.width = 9, 16
        _poner(ws, ch, ancla)


def _hoja_costes(ws, res):
    _anchos(ws, [44, 16, 4, 22, 14])
    ws["A1"] = "Desglose de costes y de inversión"
    ws["A1"].font = TIT
    _cab(ws, 3, ["Coste anual", "€/año"])
    _tabla(ws, 4, [("Electricidad (sistema + enfriadora residual)", res["c_elec"], "#,##0"),
                   ("Agua", res["c_agua"], "#,##0"), ("Mantenimiento", res["c_mant"], "#,##0"),
                   ("CAPEX anualizado (CRF)", res["capex_anual"], "#,##0"),
                   ("Crédito por calor desplazado (−)", -res["credito_calor"], "#,##0"),
                   ("TOTAL J", res["J"], "#,##0"), ("Línea base", res["J_base"], "#,##0"),
                   ("AHORRO NETO", res["ahorro_neto"], "#,##0")])
    _cab(ws, 3, ["Inversión", "€"], 4)
    inv = [("Absorción", res["capex_abs"]), ("BdC 1", res["capex_bdc1"]), ("Adiabático", res["capex_adiab"]),
           ("R1 intercambiador", res["capex_r1"]), ("BdC R2", res["capex_r2"]), ("BoP", res["capex_bop"])]
    for i, (n, v) in enumerate(inv):
        ws.cell(4 + i, 4, n)
        ws.cell(4 + i, 5, v).number_format = "#,##0"
    ws.cell(10, 4, "TOTAL").font = Font(bold=True)
    ws.cell(10, 5, res["capex_total"]).number_format = "#,##0"
    ch = BarChart()
    ch.type, ch.title, ch.legend = "bar", "Reparto de la inversión (€)", None
    ch.add_data(Reference(ws, min_col=5, min_row=4, max_row=9), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=4, min_row=4, max_row=9))
    ch.height, ch.width = 8, 14
    _poner(ws, ch, "D13")
    ws["A14"] = "Electricidad por consumo (MWh/a)"
    ws["A14"].font = SUB
    _tabla(ws, 15, [("BdC 1", res["e_bdc1_mwh"], "0.0"), ("BdC R2", res["e_r2_mwh"], "0.0"),
                    ("Bombas", res["e_bombas_mwh"], "0.0"), ("Ventiladores", res["e_vent_mwh"], "0.0"),
                    ("Auxiliares absorción", res["e_aux_mwh"], "0.0"),
                    ("Ahorro dry cooler CPD (−)", -res["e_dc_ahorro_mwh"], "0.0"),
                    ("Enfriadora residual", res["elec_enf_mwh"], "0.0"),
                    ("Línea base (solo enfriadora)", res["elec_base_mwh"], "0.0")])


def _hoja_sensibilidad(ws, res, sens):
    _anchos(ws, [46, 12, 16, 16, 14, 14])
    ws["A1"] = "Sensibilidad ±20 % de cada parámetro sobre el ahorro neto del caso"
    ws["A1"].font = TIT
    _cab(ws, 3, ["Parámetro", "Valor", "Ahorro con −20 %", "Ahorro con +20 %", "Δ con −20 %", "Δ con +20 %"])
    for i, s in enumerate(sens):
        for j, v in enumerate([s["parametro"], s["valor"], s["ahorro_menos"], s["ahorro_mas"],
                               s["ahorro_menos"] - s["ahorro_ref"], s["ahorro_mas"] - s["ahorro_ref"]]):
            c = ws.cell(4 + i, 1 + j, v)
            if j >= 2:
                c.number_format = "#,##0"
    n = len(sens)
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap = "bar", "clustered", 100
    ch.title = f"Tornado: cambio del ahorro neto (ref. {res['ahorro_neto']:,.0f} €/a)"
    ch.x_axis.scaling.orientation = "maxMin"  # el parámetro más influyente arriba
    ch.y_axis.crosses = "max"                 # y el eje de euros abajo
    ch.add_data(Reference(ws, min_col=5, max_col=6, min_row=3, max_row=3 + n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3 + n))
    ch.height, ch.width = 11, 18
    _poner(ws, ch, f"A{n + 6}")


def avisos(modo, caso, res, filas):
    out = []
    eqs = [("BdC 1", "cal_bdc1", caso["eq.bdc1"]), ("Absorción", "cal_abs", caso["eq.abs"])]
    if modo == "con_recuperacion":
        eqs.append(("BdC R2", "cal_r2", caso["eq.r2"]))
    for n_eq, k, m in eqs:
        if res[k] >= 2:
            out.append(f"{EMOJI_CAL[res[k]]}: {n_eq} '{m}' trabaja fuera de las condiciones publicadas por el fabricante. "
                       "Pide la curva completa antes de citar su COP.")
        elif res[k] == 0:
            out.append(f"{EMOJI_CAL[0]}: {n_eq} es un equipo genérico (un solo punto, capacidad libre), no un modelo real.")
    if res["cobertura_frio"] < 0.999:
        out.append(f"A plena carga la absorción cubre el {res['cobertura_frio']:.0%} de la demanda; el resto lo pone la "
                   f"enfriadora existente. Limita: {res['limitante']}.")
    if res["nom_adiab_kw"] > 0:
        out.append(f"El rechazo supera la torre existente: hace falta un adiabático de {res['nom_adiab_kw']:.0f} kW reales.")
    if filas:
        malos = sum(1 for f in filas if max(f["cal_bdc1"], f["cal_abs"], f["cal_r2"]) == 3)
        if malos:
            out.append(f"{malos} de {len(filas)} escenarios del barrido usan algún equipo fuera de rango (>10 K). "
                       "Están marcados en rojo en la hoja Barrido.")
    out += ["Clima fijo: el rechazo se calcula a las T de la hoja Entradas todo el año (sin clima horario).",
            "Los importes en euros dependen de los parámetros [SUPUESTO] de la hoja Entradas (CAPEX, precios).",
            "Línea base: toda la demanda de frío con la enfriadora existente (ya amortizada)."]
    return out


def _hoja_entradas(ws, caso, barrido_vars):
    _anchos(ws, [34, 52, 40, 18, 11, 22])
    _cab(ws, 1, ["Bloque", "Parámetro", "Valor", "Unidad", "Origen", "ruta"])
    for bloque, ruta_p, etq, _, uni, origen, _ in P.PARAMS:
        ws.append([bloque, etq, caso[ruta_p], uni, origen, ruta_p])
    for (ruta_e, etq, _, _), (ruta_n, etq_n, _) in zip(P.EQUIPOS, P.UNIDADES):
        ws.append(["Equipos", etq, caso[ruta_e], "", "", ruta_e])
        ws.append(["Equipos", etq_n, caso[ruta_n], "0 = automático", "", ruta_n])
    for carga, peso in caso["perfil"]:
        ws.append(["Perfil", f"Carga {carga:.0%}", peso, "fracción de horas", "", "perfil"])
    for ruta_b, vals in barrido_vars:
        ws.append(["Barrido", P.etiqueta(ruta_b), "; ".join(f"{v:g}" for v in vals), "", "", ruta_b])


def escribir(ruta, modo, caso, res, barrido_vars, filas_barrido, sens, equil, cat):
    wb = Workbook()
    wb.active.title = "Resumen"  # antes de crear gráficos: sus referencias llevan el nombre de la hoja
    _hoja_resumen(wb.active, modo, caso, res, equil)
    _hoja_balance(wb.create_sheet("Balance"), caso, res["diseno"])
    _hoja_equipos(wb.create_sheet("Equipos"), modo, caso, res, cat)
    if filas_barrido:
        rutas = [v[0] for v in barrido_vars]
        _hoja_barrido(wb, rutas, filas_barrido)
        _hoja_mapas(wb, rutas, filas_barrido)
    _hoja_anual(wb.create_sheet("Anual"), res)
    _hoja_costes(wb.create_sheet("Costes"), res)
    _hoja_sensibilidad(wb.create_sheet("Sensibilidad"), res, sens)
    ws = wb.create_sheet("Avisos")
    _anchos(ws, [120])
    ws["A1"] = "Avisos y limitaciones"
    ws["A1"].font = TIT
    for i, a in enumerate(avisos(modo, caso, res, filas_barrido)):
        ws.cell(3 + i, 1, "• " + a).alignment = Alignment(wrap_text=True)
    _hoja_entradas(wb.create_sheet("Entradas"), caso, barrido_vars)
    Path(ruta).parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
