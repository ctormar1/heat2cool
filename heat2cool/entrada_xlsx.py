"""Plantilla entrada.xlsx: la escribe con los valores por defecto y la lee de vuelta como un caso."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from . import parametros as P

AZUL = PatternFill("solid", fgColor="DDEBF7")
GRIS = PatternFill("solid", fgColor="F2F2F2")
CAB = PatternFill("solid", fgColor="1F4E78")
BLANCO = Font(color="FFFFFF", bold=True)


def _cabecera(ws, cols, anchos):
    ws.append(cols)
    for i, w in enumerate(anchos, 1):
        ws.cell(1, i).fill, ws.cell(1, i).font = CAB, BLANCO
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"


def escribir_plantilla(ruta: str | Path, cat: dict, caso: dict | None = None, barrido=None):
    caso = caso or P.caso_por_defecto()
    barrido = barrido or P.BARRIDO
    wb = Workbook()

    ws = wb.active
    ws.title = "Entradas"
    _cabecera(ws, ["Bloque", "Parámetro", "Valor", "Unidad", "Origen", "Nota", "ruta (no tocar)"],
              [34, 52, 12, 18, 11, 50, 22])
    for bloque, ruta_p, etq, _, uni, origen, nota in P.PARAMS:
        ws.append([bloque, etq, caso[ruta_p], uni, origen, nota, ruta_p])
        r = ws.max_row
        ws.cell(r, 3).fill = AZUL
        if origen == P.SUP:
            ws.cell(r, 5).font = Font(color="C65911", bold=True)
        ws.cell(r, 7).font = Font(color="A6A6A6")

    ws = wb.create_sheet("Equipos")
    _cabecera(ws, ["Equipo", "Modelo (elige del desplegable)", "Nº unidades (0 = automático)", "ruta (no tocar)",
                   "ruta n (no tocar)"], [40, 60, 26, 14, 14])
    hc = wb.create_sheet("Catálogo")
    _cabecera(hc, ["Bombas de calor", "Absorción"], [60, 60])
    bdc = [n for n, e in cat.items() if e["tipo"] == "bdc"]
    ab = [n for n, e in cat.items() if e["tipo"] == "absorcion"]
    for i in range(max(len(bdc), len(ab))):
        hc.append([bdc[i] if i < len(bdc) else None, ab[i] if i < len(ab) else None])
    listas = {"bdc": f"=Catálogo!$A$2:$A${len(bdc) + 1}", "absorcion": f"=Catálogo!$B$2:$B${len(ab) + 1}"}
    for (ruta_e, etq, tipo, _), (ruta_n, _, _) in zip(P.EQUIPOS, P.UNIDADES):
        ws.append([etq, caso[ruta_e], caso[ruta_n], ruta_e, ruta_n])
        r = ws.max_row
        dv = DataValidation(type="list", formula1=listas[tipo], allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(ws.cell(r, 2))
        ws.cell(r, 2).fill = ws.cell(r, 3).fill = AZUL
    ws.append([])
    ws.append(["El catálogo completo (puntos de cada curva) está en catalogo/curvas.csv. Añade filas allí para meter curvas nuevas."])

    ws = wb.create_sheet("Perfil")
    _cabecera(ws, ["Carga del CPD (fracción)", "Fracción de horas del año"], [26, 26])
    for carga, peso in caso["perfil"]:
        ws.append([carga, peso])
        ws.cell(ws.max_row, 1).fill = ws.cell(ws.max_row, 2).fill = AZUL
    ws.append([])
    ws.append(["Las fracciones de horas deben sumar 1. Puedes añadir o quitar filas."])

    ws = wb.create_sheet("Barrido")
    _cabecera(ws, ["ruta del parámetro a barrer", "valores separados por ;", "Parámetro"], [26, 40, 52])
    for ruta_b, vals in barrido:
        ws.append([ruta_b, "; ".join(f"{v:g}" for v in vals), P.etiqueta(ruta_b)])
        ws.cell(ws.max_row, 1).fill = ws.cell(ws.max_row, 2).fill = AZUL
    ws.append([])
    ws.append(["Pon cualquier 'ruta' de la hoja Entradas (columna G). Se calculan todas las combinaciones."])
    ws.append(["Deja la hoja sin filas para calcular solo el caso de la hoja Entradas."])
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="center")

    Path(ruta).parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)


def leer(ruta: str | Path):
    """Devuelve (caso, barrido) a partir del Excel de entrada."""
    wb = load_workbook(ruta, data_only=True)
    caso = P.caso_por_defecto()
    validas = {p[1] for p in P.PARAMS}
    for row in wb["Entradas"].iter_rows(min_row=2, values_only=True):
        ruta_p, valor = row[6], row[2]
        if ruta_p in validas:
            if valor is None or isinstance(valor, str):
                raise ValueError(f"Hoja Entradas: '{row[1]}' ({ruta_p}) necesita un número, tiene {valor!r}")
            caso[ruta_p] = float(valor)
    for row in wb["Equipos"].iter_rows(min_row=2, values_only=True):
        if row[3] and str(row[3]).startswith("eq."):
            caso[row[3]] = str(row[1]).strip()
            caso[row[4]] = int(row[2] or 0)
    perfil = [[float(a), float(b)] for a, b, *_ in wb["Perfil"].iter_rows(min_row=2, values_only=True)
              if isinstance(a, (int, float)) and isinstance(b, (int, float))]
    s = sum(b for _, b in perfil)
    if abs(s - 1.0) > 1e-3:
        raise ValueError(f"Hoja Perfil: las fracciones de horas suman {s:.3f}, deben sumar 1")
    caso["perfil"] = perfil
    barrido = []
    for row in wb["Barrido"].iter_rows(min_row=2, values_only=True):
        if row[0] in validas and row[1] is not None:
            vals = [float(x.replace(",", ".")) for x in str(row[1]).split(";") if x.strip()]
            barrido.append((row[0], vals))
    return caso, barrido
